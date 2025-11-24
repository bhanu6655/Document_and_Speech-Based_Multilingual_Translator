# app.py

import os
import tkinter as tk
from tkinter import END, filedialog, messagebox
from tkinter import ttk

from docx import Document

from config import LANG_NAMES, LANG_NAME_TO_CODE
from translation_engine import translate_text
from audio_engine import (
    recognize_from_audio_file,
    recognize_from_microphone,
    speak_text,
    stop_playback,
    save_tts_to_file,
    cleanup_audio,
)


def create_button(parent, text, command, row, column):
    btn = tk.Button(
        parent,
        text=text,
        command=command,
        font=("Arial", 12, "bold"),
        bg="#3498db",
        fg="white",
        activebackground="#2980b9",
        activeforeground="white",
        padx=15,
        pady=8,
        bd=0,
        relief="ridge",
        cursor="hand2",
    )
    btn.grid(row=row, column=column, padx=10, pady=10, sticky="ew")
    return btn


def main():
    # === Main window ===
    main_win = tk.Tk()
    main_win.title("Advanced Multilingual Translation Platform")
    main_win.geometry("1000x750")
    main_win.config(bg="#1E1E1E")
    main_win.resizable(False, False)

    # === Language selection variables ===
    v1 = tk.StringVar(main_win)
    v1.set(LANG_NAMES[0])   # Input language default: English

    v2 = tk.StringVar(main_win)
    v2.set(LANG_NAMES[1])   # Output language default: Hindi

    # === Title ===
    ttk.Label(
        main_win,
        text="Advanced Multilingual Translation Platform",
        font=("Arial", 16, "bold"),
        background="#1E1E1E",
        foreground="white",
    ).pack(pady=10)

    # === Text frames ===
    frame1 = tk.Frame(main_win, bg="#333333", padx=10, pady=10)
    frame1.pack(pady=10, fill="x")

    ttk.Label(frame1, text="Input Text:", background="#333333", foreground="white").grid(
        row=0, column=0, padx=10, pady=5
    )
    input_text = tk.Text(frame1, height=10, width=50, font=("Arial", 12))
    input_text.grid(row=1, column=0, padx=10, pady=5)

    ttk.Label(frame1, text="Translated Text:", background="#333333", foreground="white").grid(
        row=0, column=1, padx=10, pady=5
    )
    output_text = tk.Text(frame1, height=10, width=50, font=("Arial", 12))
    output_text.grid(row=1, column=1, padx=10, pady=5)

    # === Button callbacks using engines ===

    def upload_text_file():
        file_path = filedialog.askopenfilename(
            filetypes=[("Text/DOCX Files", "*.txt *.docx"), ("All Files", "*.*")]
        )
        if not file_path:
            return

        input_text.delete("1.0", END)

        try:
            if file_path.lower().endswith(".txt"):
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    input_text.insert(END, f.read())

            elif file_path.lower().endswith(".docx"):
                doc = Document(file_path)
                content = [p.text for p in doc.paragraphs if p.text.strip()]
                # Include table texts
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                content.append(cell.text)

                input_text.insert(END, "\n".join(content))

            else:
                messagebox.showerror("Error", "Unsupported file format")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    def upload_audio_file():
        file_path = filedialog.askopenfilename(
            filetypes=[("Audio Files", "*.wav *.mp3 *.ogg *.flac"), ("All Files", "*.*")]
        )
        if not file_path:
            return

        input_text.delete("1.0", END)

        try:
            recognized = recognize_from_audio_file(file_path, v1.get())
            input_text.insert(END, recognized)
        except Exception as e:
            messagebox.showerror("Audio Error", str(e))

    def listen_and_translate():
        input_text.delete("1.0", END)
        output_text.delete("1.0", END)

        try:
            # Inform the user you're listening
            messagebox.showinfo("Speak Now", "Listening... Speak clearly into the microphone")

            spoken_text = recognize_from_microphone(v1.get())
            input_text.insert(END, spoken_text)

            translated = translate_text(spoken_text, v2.get())
            output_text.insert(END, translated)
        except Exception as e:
            messagebox.showerror("Speech Error", str(e))

    def do_translate():
        output_text.delete("1.0", END)
        text = input_text.get("1.0", END).strip()
        if not text:
            messagebox.showwarning("Warning", "Please enter text to translate")
            return

        try:
            translated = translate_text(text, v2.get())
            output_text.insert(END, translated)
        except Exception as e:
            messagebox.showerror("Translation Error", str(e))

    def do_speak():
        text = output_text.get("1.0", END).strip()
        if not text:
            messagebox.showwarning("Warning", "No translated text to speak")
            return

        try:
            speak_text(text, v2.get())
        except Exception as e:
            messagebox.showerror("Speech Error", str(e))

    def do_stop_speaking():
        stop_playback()

    def save_translated_text():
        text = output_text.get("1.0", END).strip()
        if not text:
            messagebox.showwarning("Warning", "No translated text to save")
            return

        file_path = filedialog.asksaveasfilename(
            defaultextension=".txt",
            filetypes=[("Text files", "*.txt"), ("Word Documents", "*.docx")],
        )
        if not file_path:
            return

        try:
            if file_path.lower().endswith(".docx"):
                doc = Document()
                for line in text.splitlines():
                    doc.add_paragraph(line)
                doc.save(file_path)
            else:
                with open(file_path, "w", encoding="utf-8") as f:
                    f.write(text)

            messagebox.showinfo("Saved", "File saved successfully")
        except Exception as e:
            messagebox.showerror("Save Error", str(e))

    def save_voice():
        text = output_text.get("1.0", END).strip()
        if not text:
            messagebox.showwarning("Warning", "No text to convert")
            return

        file_path = filedialog.asksaveasfilename(
            defaultextension=".mp3",
            filetypes=[("MP3 files", "*.mp3"), ("WAV files", "*.wav")],
        )
        if not file_path:
            return

        try:
            save_tts_to_file(text, v2.get(), file_path)
            messagebox.showinfo("Success", "Audio file saved")
        except Exception as e:
            messagebox.showerror("Audio Save Error", str(e))

    # === Buttons ===
    button_frame = tk.Frame(main_win, bg="#1E1E1E")
    button_frame.pack(pady=10)

    for i in range(3):
        button_frame.grid_columnconfigure(i, weight=1)

    buttons = [
        ("\U0001F4C4 Upload File (TXT/DOCX)", upload_text_file),
        ("\U0001F3A4 Upload Audio File", upload_audio_file),
        ("\U0001F3A7 Listen and Translate", listen_and_translate),
        ("\U0001F4DD Translate", do_translate),
        ("\U0001F50A Speak", do_speak),
        ("\U0001F507 Stop Speaking", do_stop_speaking),
        ("\U0001F4BE Save Text", save_translated_text),
        ("\U0001F3B5 Save Audio", save_voice),
    ]

    for index, (text, cmd) in enumerate(buttons):
        create_button(button_frame, text, cmd, index // 3, index % 3)

    # === Language selectors ===
    lang_frame = tk.Frame(main_win, bg="#1E1E1E")
    lang_frame.pack(pady=10)

    ttk.Label(
        lang_frame,
        text="Select Input Audio Language:",
        background="#1E1E1E",
        foreground="white",
    ).pack(side="left")
    ttk.Combobox(
        lang_frame,
        textvariable=v1,
        values=LANG_NAMES,
        state="readonly",
        width=20,
    ).pack(side="left", padx=10)

    ttk.Label(
        lang_frame,
        text="Select Output Language:",
        background="#1E1E1E",
        foreground="white",
    ).pack(side="left")
    ttk.Combobox(
        lang_frame,
        textvariable=v2,
        values=LANG_NAMES,
        state="readonly",
        width=20,
    ).pack(side="left", padx=10)

    # === Handle proper cleanup on close ===
    def on_close():
        cleanup_audio()
        main_win.destroy()

    main_win.protocol("WM_DELETE_WINDOW", on_close)

    # Start the Tkinter event loop
    main_win.mainloop()


if __name__ == "__main__":
    main()
